from pathlib import Path
import re,json,subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path('/workspace/work/3cx')
speech=(root/'speech_draft.md').read_text()
parts={int(n):[title,body.strip()] for n,title,body in re.findall(r'^## (\d+)\. ([^\n]+)\n\n(.*?)(?=^## |\Z)',speech,re.M|re.S)}
assert len(parts)==14
def spoken(body):
    no_stage=re.sub(r'^\[[^\n]+\]\s*$', '',body,flags=re.M)
    return sum(bool(re.search(r'[A-Za-zА-Яа-яЁё0-9]',t)) for t in no_stage.split())
counts={n:spoken(body) for n,(_,body) in parts.items()}
assert 2400<=sum(counts.values())<=2900
rate=(sum(counts.values())-counts[7])/22.5
extras={1:5,2:2,3:3,4:8,5:3,6:3,8:8,9:3,10:2,11:8,12:2,13:3}
dur={n:100 if n==7 else round((c/rate*60+extras.get(n,0))/5)*5 for n,c in counts.items()}
dur[14]+=1500-sum(dur.values())
def clock(t):return f'{t//60:02}:{t%60:02}'
times={};t=0
for n in range(1,15):
    times[n]=clock(t)+'–'+clock(t+dur[n]);t+=dur[n]
assert t==1500
sources={2:['C1','C3','C6','C5'],3:['C1','C4'],4:['N1','N2'],5:['C2','D3','D4','D6','N2'],6:['D5'],7:['C2','D3','D4'],8:['C3','C4'],9:['C5','D1','D2'],10:['C5']}
links={
'C1':('Sophos','https://www.sophos.com/en-gb/blog/3cx-dll-sideloading-attack'),
'C2':('Sophos CSV','https://raw.githubusercontent.com/sophoslabs/IoCs/master/3CX%20IoCs%202023-03.csv'),
'C3':('SentinelOne','https://www.sentinelone.com/blog/smoothoperator-ongoing-campaign-trojanizes-3cx-software-in-software-supply-chain-attack/'),
'C4':('Unit 42','https://unit42.paloaltonetworks.com/3cxdesktopapp-supply-chain-attack/'),
'C5':('Mandiant','https://cloud.google.com/blog/topics/threat-intelligence/3cx-software-supply-chain-compromise'),
'C6':('3CX','https://www.3cx.com/blog/news/desktopapp-security-alert/'),
'D1':('AbuseIPDB CHECK','https://docs.abuseipdb.com/#check-endpoint'),
'D2':('AbuseIPDB FAQ','https://www.abuseipdb.com/faq.html'),
'D3':('VT Analysis','https://docs.virustotal.com/reference/analyses-object'),
'D4':('VT File object','https://docs.virustotal.com/reference/files'),
'D5':('VT Searching','https://docs.virustotal.com/docs/searching'),
'D6':('VT false positives','https://docs.virustotal.com/docs/virustotal-intelligence-introduction'),
'N1':('NIST IoC','https://csrc.nist.gov/glossary/term/indicator_of_compromise'),
'N2':('Microsoft Authenticode / хеш','https://learn.microsoft.com/en-us/windows-hardware/drivers/install/authenticode')}
blocks=['## 4. Полный текст устного выступления',f'Произносимая часть: **{sum(counts.values())} слов**, из них {counts[7]} — реплики демонстрации. Ориентир: около **{rate:.0f} слов в минуту** вне демонстрации, **1:40** на демонстрацию и примерно **50 секунд** на короткие паузы и переходы. Это расчёт, фактическая длительность определяется репетицией. Заголовки, пометы в квадратных скобках и редакторские основания не произносятся.']
for n,(title,body) in parts.items():
    blocks.append(f'### Слайд {n}. {title} • {times[n]}\n\n{body}')
    if n in sources:
        selected=dict(links)
        if n==4:selected['N2']=('Microsoft Get-FileHash','https://learn.microsoft.com/en-us/powershell/module/microsoft.powershell.utility/get-filehash')
        citations=', '.join(f'[{i}: {selected[i][0]}]({selected[i][1]})' for i in sources[n])
        blocks.append(f'*Редакторское основание, не произносить: {citations}.*')
table='| Время | Слайды | Вопрос блока |\n| --- | --- | --- |\n'
groups=[(1,2,'Что известно об исходном событии и реальном кейсе?'),(3,4,'Какой компонент и какие точные байты проверяем?'),(5,6,'Что говорит отчёт и что раскрывает запрос?'),(7,7,'Короткий разбор опубликованного хеша.'),(8,9,'Как перейти к сетевому объекту, сохранив время и основание?'),(10,12,'Что произошло на нашем устройстве и как записать вывод?'),(13,14,'Что искать в истории и какой следующий шаг?')]
for a,b,label in groups:
    interval=times[a].split('–')[0]+'–'+times[b].split('–')[1]
    table+=f'| {interval} | {a if a==b else str(a)+"–"+str(b)} | {label} |\n'
slides=(root/'slides.md').read_text()
slides=re.sub(r'^### (\d+)\. ([^\n]+)',lambda m:f'### {m.group(1)}. {m.group(2)} • {times[int(m.group(1))]}',slides,flags=re.M)
brief=(root/'brief.md').read_text()
brief=brief.replace('<!-- TIMING -->',table.strip()).replace('<!-- SPEECH -->','\n\n'.join(blocks)).replace('<!-- SLIDES -->',slides)
assert '<!--' not in brief
brief=brief.replace('Первое чтение — с секундомером.',f'Контрольные отметки: {times[3].split("–")[0]} (компоненты), {times[7].split("–")[0]} (демонстрация), {times[11].split("–")[0]} (внутренняя проверка), {times[13].split("–")[0]} (hunting). Первое чтение — с секундомером.')
(root/'3cx_full_talk.md').write_text(brief)
(root/'3cx_speech.md').write_text('\n\n'.join(blocks))
(root/'timing.json').write_text(json.dumps({'speech_words':sum(counts.values()),'demo_words':counts[7],'words_per_minute':rate,'counts':counts,'times':times,'durations_seconds':dur},ensure_ascii=False,indent=2))
subprocess.run(['pandoc',str(root/'3cx_full_talk.md'),'-o',str(root/'3cx_full_talk.docx'),'--reference-doc','/workspace/work/reference.docx','--resource-path',str(root)],check=True)
doc=Document(root/'3cx_full_talk.docx')
for para in doc.paragraphs:
    if para.style.name=='Heading 2':para.paragraph_format.page_break_before=True
table_weights=[[18,30,28,24],[42,58],[28,26,26,20],[20,12,68],[18,42,40],[16,42,42],[24,35,41],[7,47,46]]
section=doc.sections[0]
usable_twips=int((section.page_width-section.left_margin-section.right_margin)/635)
for tab,weights in zip(doc.tables,table_weights):
    tab.style='Table Grid'
    tab.autofit=False
    widths=[round(usable_twips*w/sum(weights)) for w in weights]
    tblw=tab._tbl.tblPr.find(qn('w:tblW'))
    tblw.set(qn('w:type'),'dxa');tblw.set(qn('w:w'),str(usable_twips))
    for col,width in zip(tab.columns,widths):col.width=Twips(width)
    for i,row in enumerate(tab.rows):
        if i==0:
            hdr=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(hdr)
        for cell,width in zip(row.cells,widths):
            cell.width=Twips(width)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E3EEF2');cell._tc.get_or_add_tcPr().append(shading)
            for para in cell.paragraphs:
                # The reference document has no Pandoc Compact style; use a defined style.
                para.style='Normal'
                para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.line_spacing=1
                para.paragraph_format.keep_with_next=False
                para.paragraph_format.keep_together=False
                for run in para.runs:
                    run.font.size=Pt(9.5)
                    if i==0:run.bold=True
for sec in doc.sections:
    footer=sec.footer.paragraphs[0];footer.text='3CX · Threat Intelligence · Черновик · ';footer.alignment=2
    for run in footer.runs:run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Загуглить противника: дело об обновлении 3CX'
doc.core_properties.subject='Самостоятельный доклад 25 минут для младших курсов: речь, 14 слайдов, демонстрация, памятка и исследование'
doc.core_properties.author='[Спикер], [Компания]'
doc.save(root/'3cx_full_talk.docx')
assert 'Слайд 14.' in '\n'.join(p.text for p in doc.paragraphs)
print(json.dumps({'speech_words':sum(counts.values()),'demo_words':counts[7],'rate':round(rate,1),'times':times,'docx_bytes':(root/'3cx_full_talk.docx').stat().st_size,'tables':len(doc.tables)},ensure_ascii=False))
