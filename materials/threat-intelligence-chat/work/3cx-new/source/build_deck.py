from pathlib import Path
import json,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parent
P=Presentation();P.slide_width=Inches(13.333333);P.slide_height=Inches(7.5)
P.core_properties.title='Загуглить противника: думать с помощью TI'
P.core_properties.subject='3CX · 8 слайдов · 25 минут · полная речь в заметках'
P.core_properties.author='Компания / Имя спикера'
BG='F4F3EE';INK='172D38';TEAL='007D7A';TEAL_LIGHT='DDEDEA';MUTED='526973'
LINE='CDD7D6';WHITE='FFFFFF';DARK='142B35';DARK_CARD='223D48';CREAM='F4F3EE'
YELLOW='E1B56B';RED='B85043';ORANGE_LIGHT='F4E4D2';FONT='DejaVu Sans';MONO='DejaVu Sans Mono'
HASH='c485674ee63ec8d4e8fde9800788175a8b02d3f9416d0e763360fff7f8eb4e02'
SOURCE={
 'sophos':('Sophos · разбор 3CX, 2023','https://www.sophos.com/en-gb/blog/3cx-dll-sideloading-attack'),
 'ioc':('SophosLabs · опубликованный SHA-256','https://raw.githubusercontent.com/sophoslabs/IoCs/master/3CX%20IoCs%202023-03.csv'),
 'alert':('3CX · DesktopApp Security Alert, 30.03.2023','https://www.3cx.com/blog/news/desktopapp-security-alert/'),
 'mandiant':('Mandiant · 3CX, 20.04.2023','https://cloud.google.com/blog/topics/threat-intelligence/3cx-software-supply-chain-compromise'),
 'nist':('NIST · Indicator of Compromise','https://csrc.nist.gov/glossary/term/indicator_of_compromise'),
 'hash':('Microsoft · Get-FileHash','https://learn.microsoft.com/en-us/powershell/module/microsoft.powershell.utility/get-filehash'),
 'signature':('Microsoft · Authenticode','https://learn.microsoft.com/en-us/windows-hardware/drivers/install/authenticode'),
 'file':('VirusTotal · File object','https://docs.virustotal.com/reference/files'),
 'analysis':('VirusTotal · Analysis object','https://docs.virustotal.com/reference/analyses-object'),
 'comments':('VirusTotal · Community / Comments','https://docs.virustotal.com/docs/comments'),
 'reports':('VirusTotal · Results reports','https://docs.virustotal.com/docs/results-reports'),
 'privacy':('VirusTotal · Searching','https://docs.virustotal.com/docs/searching'),
 'abuse':('AbuseIPDB · CHECK','https://docs.abuseipdb.com/#check-endpoint'),
 'faq':('AbuseIPDB · FAQ','https://www.abuseipdb.com/faq.html'),
 'vishnyakov':('А. Вишняков · Ошибки TI в SOC, PHDays 2025','https://pt-phdays.storage.yandexcloud.net/Aleksej_Vishnyakov_Oshibki_ispolzovaniya_threat_intelligence_v_SOC_012edd864a.pdf'),
 'arefiev':('Н. Арефьев · Открытые источники TI, SOC Forum 2023','https://forumsoc.ru/upload/iblock/007/zrrjrpwyt50186gn7sg3ij6sts223qvl.pdf'),
 'nenakhov':('А. Ненахов · Как киберразведка помогает SOC, 2023','https://forumsoc.ru/upload/iblock/485/vw01piz8qd0037k4zsusyelev6851jc4.pdf'),
 'kotsyuba':('И. Коцюба · Не все индикаторы полезны, PHDays 2025','https://pt-phdays.storage.yandexcloud.net/Ilya_Koczyuba_Threat_Intelligence_ili_ne_vse_indikatory_odinakovo_polezny_7f23cf3f82.pdf'),
}
TIMES=['00:00–01:30','01:30–04:00','04:00–06:30','06:30–10:00','10:00–14:30','14:30–19:00','19:00–23:30','23:30–25:00']
TITLES=['Загуглить противника','TI помогает выбрать следующий шаг','Объект ещё нужно объяснить','Совпадение — начало проверки','Сначала вопрос. Затем сервис.','Одна иконка. Несколько компонентов.','Теперь — к своим устройствам','Что ответим сотруднику?']
SOURCES=[['alert','sophos'],['vishnyakov','nenakhov'],['nist','hash'],['ioc','signature','arefiev'],['file','analysis','comments','reports','privacy','abuse','faq'],['sophos','ioc','alert'],['sophos','nenakhov'],['vishnyakov','nenakhov']]
TRACK=[]

def rgb(v):return RGBColor.from_string(v)
def shape(sl,x,y,w,h,fill,stroke=None,kind=MSO_SHAPE.RECTANGLE,radius=.12):
    a=sl.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=rgb(fill)
    if stroke:a.line.color.rgb=rgb(stroke);a.line.width=Pt(.8)
    else:a.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:a.adjustments[0]=radius
    # Prevent the template's default shadow from changing a flat diagram.
    for ref in a._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
    a._element.spPr.append(OxmlElement('a:effectLst'))
    return a
def box(sl,x,y,w,h,fill=WHITE,stroke=LINE):return shape(sl,x,y,w,h,fill,stroke,MSO_SHAPE.ROUNDED_RECTANGLE,.06)
def tx(sl,x,y,w,h,text,size=24,color=INK,bold=False,font=FONT,align=PP_ALIGN.LEFT,link=None):
    a=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=a.text_frame
    f.clear();f.word_wrap=True;f.margin_left=f.margin_right=0;f.margin_top=f.margin_bottom=0
    f.vertical_anchor=MSO_ANCHOR.TOP
    for i,s in enumerate(str(text).split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.12
        r=p.add_run();r.text=s;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    if link:a.click_action.hyperlink.address=link
    TRACK.append({'slide':len(P.slides),'text':text,'box':[x,y,w,h],'font_size':size})
    return a
def ln(sl,x1,y1,x2,y2,color=LINE,width=1):
    a=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    a.line.color.rgb=rgb(color);a.line.width=Pt(width);return a
def arrow(sl,x,y,w=.22,h=.28,color=TEAL):return shape(sl,x,y,w,h,color,kind=MSO_SHAPE.CHEVRON)
def badge(sl,x,y,w,text,dark=False,color=TEAL):
    shape(sl,x,y,w,.34,DARK_CARD if dark else TEAL_LIGHT,kind=MSO_SHAPE.ROUNDED_RECTANGLE,radius=.2)
    tx(sl,x+.1,y+.075,w-.2,.18,text,9,color if not dark else '87D5CA',True)
def base(n,title,eyebrow,dark=False):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(DARK if dark else BG)
    tx(s,.62,.35,10.5,.25,eyebrow,10,'87D5CA' if dark else TEAL,True)
    tx(s,.62,.88,12.05,1.05,title,34,CREAM if dark else INK,True)
    ln(s,.62,6.91,12.7,6.91,'36515C' if dark else LINE,.8)
    tx(s,.62,7.08,9.5,.2,'ЗАГУГЛИТЬ ПРОТИВНИКА  /  ДУМАТЬ С ПОМОЩЬЮ TI',9,'B0C4CD' if dark else MUTED)
    tx(s,11.72,7.04,.95,.28,f'{n:02} / 08',12,'B0C4CD' if dark else MUTED,True,align=PP_ALIGN.RIGHT)
    return s
def foot(sl,n,text=None):
    if text is None:text=' · '.join(SOURCE[k][0] for k in SOURCES[n-1][:2])+' · ссылки в заметках'
    tx(sl,.65,6.58,12.0,.22,text,8.5,MUTED,link=SOURCE[SOURCES[n-1][0]][1])
def label(sl,x,y,text,w=3.5,color=TEAL):tx(sl,x,y,w,.3,text,12,color,True)
def dot(sl,x,y,d=.12,color=TEAL):return shape(sl,x,y,d,d,color,kind=MSO_SHAPE.OVAL)
def notes(sl,n):
    speech=(ROOT/'speech.md').read_text() if (ROOT/'speech.md').exists() else ''
    blocks=re.split(r'^##\s+\d+\..*\n',speech,flags=re.M)[1:]
    body=blocks[n-1].strip() if len(blocks)==8 else '[Текст речи ещё не присоединён.]'
    note=f'СЛАЙД {n}. {TITLES[n-1]}\nТАЙМИНГ: {TIMES[n-1]}\n\n'+body
    note+='\n\nИСТОЧНИКИ И РЕДАКТОРСКИЕ ОСНОВАНИЯ:\n'+'\n\n'.join(SOURCE[k][0]+'\n'+SOURCE[k][1] for k in SOURCES[n-1])
    if n==5:note+='\n\nДЕМО: только публичный объект; готовый отчёт отличается от Upload/Reanalyze. Поиск сам по себе не обещает приватность. Текущая Community выбранного хеша не получена. AbuseIPDB CHECK — последние1–365дней, не архив2023. Полный запасной сценарий — в сопроводительных материалах.'
    if n==6:note+='\n\nПУБЛИЧНЫЙ SHA-256:\n'+HASH+'\nОПЦИОНАЛЬНЫЙ ГОТОВЫЙ ОТЧЁТ:\nhttps://www.virustotal.com/gui/file/'+HASH+'/detection\n\nКадры основаны на опубликованных материалах2023. Текущие детекты/даты анализа/комментарии не заявляются. Учебное устройство не исследовано.'
    note+='\n\nПроверка документации:21 октября2026. Подготовленные кадры с аналитическими схемами не являются скриншотами текущих карточек сервисов.'
    sl.notes_slide.notes_text_frame.text=note

# 01: the one case opens here.
s=base(1,'Загуглить\nпротивника','THREAT INTELLIGENCE ДЛЯ НАЧИНАЮЩИХ',True)
# Replace cover headline with a larger title, preserving a roomy right column.
for a in list(s.shapes):
    if a.has_text_frame and a.text=='Загуглить\nпротивника':a._element.getparent().remove(a._element)
tx(s,.62,1.12,7.15,2.0,'Загуглить\nпротивника',54,CREAM,True)
tx(s,.67,3.66,6.5,1.12,'Как думать с помощью TI\nна примере 3CX',26,CREAM)
tx(s,.67,5.38,6.25,.45,'Компания  /  Имя спикера',14,'B0C4CD')
box(s,8.0,1.45,4.66,4.58,DARK_CARD,'36515C')
badge(s,8.27,1.75,3.28,'УЧЕБНАЯ СИТУАЦИЯ · МАРТ 2023',True)
tx(s,8.28,2.44,4.05,2.2,'После обновления\nзащита предупредила.\n\nПрограмма подписана.\nМожно продолжать?',22,CREAM)
ln(s,8.29,5.12,12.32,5.12,'4B6872')
tx(s,8.29,5.38,4.0,.3,'Откроем улики по одной.',16,'87D5CA',True)
notes(s,1)

# 02: intelligence includes analysis and an internal question.
s=base(2,TITLES[1],'01 / ЗАЧЕМ НУЖЕН TI')
tx(s,.67,2.04,11.8,.85,'Сведения об угрозах помогают\nпринимать решения по защите.',27,INK)
cards=[('ВНЕШНИЕ СВЕДЕНИЯ','Что известно\nисследователям?'),('ОЦЕНКА ОСНОВАНИЙ','Что применимо\nк нашему событию?'),('СВОИ ДАННЫЕ','Чем это\nподтвердить у нас?')]
for i,(h,b) in enumerate(cards):
    x=.67+i*4.12;box(s,x,3.44,3.75,1.75)
    label(s,x+.22,3.68,h,3.35)
    tx(s,x+.22,4.18,3.35,.87,b,22)
    if i<2:arrow(s,x+3.85,4.17,.16,.24)
shape(s,.67,5.58,11.99,.65,TEAL)
tx(s,.91,5.75,11.45,.38,'Сведения → оценка → следующая проверка',22,WHITE,True)
foot(s,2);notes(s,2)

# 03: observable versus contextual indicator.
s=base(3,TITLES[2],'02 / ЧТО ТАКОЕ IoC')
objects=[('IP','Сетевой адрес'),('ДОМЕН','Имя ресурса'),('URL','Ссылка с путём'),('SHA-256','Отпечаток файла')]
for i,(h,b) in enumerate(objects):
    x=.67+i*3.08;box(s,x,2.1,2.75,1.48)
    tx(s,x+.18,2.32,2.4,.42,h,25,TEAL,True,font=MONO)
    tx(s,x+.18,3.0,2.43,.35,b,17,MUTED)
tx(s,.71,4.02,11.8,.6,'IoC — признак с основанием подозревать атаку.',27,INK,True)
ln(s,.72,4.94,12.63,4.94)
tx(s,.75,5.23,4.9,.8,'Адрес в журнале:\nнаблюдаемый объект',22,MUTED)
arrow(s,5.73,5.47,.25,.29)
tx(s,6.24,5.23,6.35,.85,'Роль + источник + время:\nоснование для проверки',22,TEAL,True)
foot(s,3);notes(s,3)

# 04: what enrichment needs to answer.
s=base(4,TITLES[3],'03 / ЧЕГО НЕ ХВАТАЕТ ДЛЯ ВЫВОДА')
questions=[('01','ТОТ ЛИ ОБЪЕКТ?','Имя или\nточные байты?'),('02','КАКАЯ РОЛЬ?','Роль объекта\nв исследовании'),('03','НА КАКУЮ ДАТУ?','Когда видели\nэту активность?'),('04','ЧТО БЫЛО У НАС?','Наличие?\nЗапуск?\nОбращение?')]
for i,(num,h,b) in enumerate(questions):
    x=.67+i*3.08;box(s,x,2.28,2.75,2.64)
    tx(s,x+.2,2.51,2.3,.45,num,27,TEAL,True)
    tx(s,x+.2,3.07,2.36,.4,h,14,INK,True)
    tx(s,x+.2,3.61,2.36,1.16,b,20,MUTED)
box(s,.67,5.32,11.99,.98,ORANGE_LIGHT,ORANGE_LIGHT)
tx(s,.93,5.5,11.48,.7,'Точный хеш может подтвердить известное содержимое.\nИспользование этого файла на устройстве проверяем отдельно.',18,INK)
foot(s,4);notes(s,4)

# 05: the services are paired with a question, no invented scores.
s=base(5,TITLES[4],'04 / КАК ПОДКРЕПИТЬ СВОЮ МЫСЛЬ ДАННЫМИ')
for x,name,sub in [(.67,'VirusTotal','Что известно об этом объекте?'),(6.79,'AbuseIPDB','Что сообщали об этом IP?')]:
    box(s,x,2.05,5.87,3.43)
    tx(s,x+.25,2.29,5.35,.5,name,28,TEAL,True)
    tx(s,x+.25,2.91,5.35,.4,sub,18,MUTED)
left=[('1','Точный объект и даты'),('2','Содержание результатов'),('3','Community → первоисточник')]
right=[('1','Период и время жалобы'),('2','Какая активность описана'),('3','Сколько разных репортеров')]
for x,items in [(.67,left),(6.79,right)]:
    for i,(n,t) in enumerate(items):
        y=3.59+i*.55;dot(s,x+.27,y+.09,.11)
        tx(s,x+.58,y,5.0,.38,t,19,INK)
tx(s,.73,5.84,11.87,.47,'Community даёт зацепку. Проверяем её по источнику.',23,INK,True)
foot(s,5,'Официальные документы VT / AbuseIPDB · Числа не равны вероятности · Полные ссылки в заметках')
notes(s,5)

# 06: a real evidence card; the teaching endpoint has no fabricated results.
s=base(6,TITLES[5],'05 / ДЕЛО ОБ ОБНОВЛЕНИИ 3CX · WINDOWS, 2023')
items=[('3CXDesktopApp.exe','Чистый\nосновной файл',WHITE,TEAL),('ffmpeg.dll','Изменённый код\nс подписью 3CX',ORANGE_LIGHT,RED),('d3dcompiler_47.dll','Добавленная\nскрытая нагрузка',WHITE,TEAL)]
for i,(h,b,col,accent) in enumerate(items):
    x=.67+i*3.08;box(s,x,2.12,2.75,2.18,col,col if col!=WHITE else LINE)
    # A minimal editable file icon.
    shape(s,x+.2,2.35,.3,.4,col,accent)
    tx(s,x+.2,2.99,2.36,.5,h,12.2,INK,True,font=MONO)
    tx(s,x+.2,3.54,2.36,.62,b,16,MUTED)
box(s,10.0,2.12,2.66,2.18,TEAL_LIGHT,TEAL_LIGHT)
tx(s,10.22,2.36,2.21,.28,'СЛЕДУЮЩИЙ ВОПРОС',10,TEAL,True)
tx(s,10.22,2.95,2.21,1.3,'Где этот\nкомпонент\nбыл у нас?',24,TEAL,True)
label(s,.73,4.78,'ОПУБЛИКОВАННЫЙ SHA-256 КОНКРЕТНОГО ffmpeg.dll',11.8)
box(s,.67,5.24,11.99,.63,DARK,DARK)
tx(s,.88,5.44,11.58,.28,HASH,14,CREAM,font=MONO,link='https://www.virustotal.com/gui/file/'+HASH+'/detection')
tx(s,.73,6.04,11.9,.29,'Исторические сведения Sophos. Наше учебное устройство ещё не проверено.',16,MUTED)
foot(s,6);notes(s,6)

# 07: local checks and the broader hypothesis.
s=base(7,TITLES[6],'06 / ОТ ОБОГАЩЕНИЯ — К ПРОВЕРКЕ')
stages=[('ПРИСУТСТВОВАЛ','Версия · путь · хеш'),('ИСПОЛНЯЛСЯ','Запуск · загрузка DLL'),('ОБРАЩАЛСЯ','Направление · процесс · время')]
for i,(h,b) in enumerate(stages):
    x=.67+i*4.12;box(s,x,2.16,3.75,1.61)
    tx(s,x+.21,2.46,3.33,.37,h,18,TEAL,True)
    tx(s,x+.21,3.11,3.33,.58,b,15.5,MUTED)
    if i<2:arrow(s,x+3.85,2.9,.16,.24)
box(s,.67,4.18,4.15,1.73)
label(s,.92,4.44,'ПОИСК ИЗВЕСТНОГО IoC',3.65)
tx(s,.92,5.0,3.65,.85,'Где встречался\nэтот точный хеш?',22,INK,True)
box(s,5.06,4.18,7.6,1.73,TEAL_LIGHT,TEAL_LIGHT)
label(s,5.32,4.44,'HUNTING-ГИПОТЕЗА',6.99)
tx(s,5.32,5.0,6.99,.85,'Что изменилось в поведении клиента\nпосле обновления?',22,TEAL,True)
tx(s,.75,6.12,11.8,.26,'Архивное упражнение: март–апрель 2023, при наличии данных. Кандидат требует проверки.',14.5,MUTED)
foot(s,7);notes(s,7)

# 08: return to the employee and leave a usable method.
s=base(8,TITLES[7],'07 / ВОЗВРАЩАЕМСЯ К ИСХОДНОМУ СОБЫТИЮ',True)
tx(s,.7,2.11,6.32,1.33,'«Программа подписана.\nМожно продолжать?»',30,CREAM,True)
for i,t in enumerate(['Уточним предупреждение и компонент.','Сверим источник, роль и время.','Проверим историю устройства.']):
    y=2.08+i*.87;dot(s,7.43,y+.14,.12,'87D5CA');tx(s,7.78,y,4.76,.84,t,21,CREAM)
steps=['Объект','Вопрос','Источник','Проверка','Действие']
for i,t in enumerate(steps):
    x=.7+i*2.43;box(s,x,4.93,2.13,.65,DARK_CARD,DARK_CARD)
    tx(s,x+.06,5.1,2.01,.3,t,18,'87D5CA',True,align=PP_ALIGN.CENTER)
    if i<4:arrow(s,x+2.22,5.13,.13,.21,'87D5CA')
tx(s,.73,6.02,11.85,.43,'Спасибо. Вопросы?',23,CREAM,True)
notes(s,8)

assert len(P.slides)==8
for item in TRACK:
    x,y,w,h=item['box'];assert min(x,y,w,h)>=0 and x+w<=13.34 and y+h<=7.5,item
P.save(ROOT/'3cx_ti_presentation.pptx')
(ROOT/'layout.json').write_text(json.dumps(TRACK,ensure_ascii=False,indent=2))
(ROOT/'sources.json').write_text(json.dumps(SOURCE,ensure_ascii=False,indent=2))
(ROOT/'slides.json').write_text(json.dumps([{'number':i+1,'title':TITLES[i],'timing':TIMES[i],'image':f'slides/{i+1:02}.png'} for i in range(8)],ensure_ascii=False,indent=2))
print(json.dumps({'slides':len(P.slides),'text_boxes':len(TRACK),'pptx_bytes':(ROOT/'3cx_ti_presentation.pptx').stat().st_size},ensure_ascii=False))
