from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import subprocess

work = Path('/workspace/work')
ref = Document()
sec = ref.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
for name in ['Normal', 'Body Text']:
    style = ref.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.12
for level, size in [(1,22),(2,17),(3,13)]:
    style = ref.styles[f'Heading {level}']
    style.font.name, style.font.size = 'Calibri', Pt(size)
    style.font.color.rgb = RGBColor.from_string('17556A')
    style.paragraph_format.keep_with_next = True
ref.save(work/'reference.docx')
subprocess.run(['pandoc',str(work/'ti_talk_full.md'),'-o',str(work/'ti_talk_full.docx'),'--reference-doc',str(work/'reference.docx')],check=True)
doc = Document(work/'ti_talk_full.docx')
for para in doc.paragraphs:
    if para.style.name == 'Heading 2':
        para.paragraph_format.page_break_before = True
for table in doc.tables:
    table.style = 'Light Shading Accent 1'
    for ri,row in enumerate(table.rows):
        pr = row._tr.get_or_add_trPr()
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            pr.append(repeat)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(9.5)
for sec in doc.sections:
    footer=sec.footer.paragraphs[0]
    footer.alignment=2
    r=footer.add_run('Threat Intelligence · Черновик · ')
    r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Загуглить противника: Threat Intelligence для начинающих'
doc.core_properties.subject = '25 минут: речь, бриф слайдов, демонстрация, памятка и источники'
doc.core_properties.author = '[Имя спикера], [Название компании]'
doc.save(work/'ti_talk_full.docx')
text = '\n'.join(p.text for p in doc.paragraphs)
assert 'Полный текст устного выступления' in text
assert 'Слайд 14.' in text
assert len(doc.tables) == 6, len(doc.tables)
print({'docx_bytes':(work/'ti_talk_full.docx').stat().st_size,'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'hyperlinks':len(doc.element.xpath('//w:hyperlink'))})
