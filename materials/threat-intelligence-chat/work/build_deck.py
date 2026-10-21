from pathlib import Path
import json, re, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path('/workspace/work')
DATA = json.loads((ROOT/'deck_copy.json').read_text())
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
P.core_properties.title = 'Загуглить противника: Threat Intelligence для начинающих'
P.core_properties.subject = '25 минут · 14 слайдов · заметки спикера · 2 приложения'
P.core_properties.author = 'Компания / Имя спикера'
BG='101B29'; CARD='1B2B3C'; TEXT='F3F6FA'; MUTED='A6B8C9'; TEAL='58D5CD'; LIME='CDF276'; ORANGE='FFB477'
PAPER='F4F6F8'; INK='142234'; DIM='526778'; LINE='DAE2E9'; WHITE='FFFFFF'
FONT='Arial'; MONO='Courier New'
COMPANY='Компания'; SPEAKER='Имя спикера'
TRACK=[]

def rgb(s): return RGBColor.from_string(s)
def rect(sl,x,y,w,h,color,line=None,radius=False):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color)
    if line:
        sh.line.color.rgb=rgb(line); sh.line.width=Pt(.8)
    else: sh.line.fill.background()
    if radius: sh.adjustments[0]=.12
    return sh

def txt(sl,x,y,w,h,string,size=24,color=TEXT,bold=False,font=FONT,align=PP_ALIGN.LEFT,link=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=Inches(.015)
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(str(string).split('\n')):
        par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        par.alignment=align; par.space_before=Pt(0); par.space_after=Pt(3)
        par.line_spacing=1.08
        run=par.add_run(); run.text=line
        run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=rgb(color)
    if link: sh.click_action.hyperlink.address=link
    TRACK.append((len(P.slides),string,x,y,w,h,size))
    return sh

def line(sl,x1,y1,x2,y2,color=TEAL,width=1.7,dash=False):
    sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def circle(sl,x,y,diam,color,linecol=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(diam),Inches(diam))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
    if linecol: sh.line.color.rgb=rgb(linecol); sh.line.width=Pt(1.5)
    else:sh.line.fill.background()
    return sh

def arrow(sl,x,y,w=.28,h=.28,color=TEAL):
    sh=sl.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh

def pill(sl,x,y,w,string,color=TEAL,dark=True):
    rect(sl,x,y,w,.35,CARD if dark else 'E5ECEF',radius=True)
    txt(sl,x+.1,y+.045,w-.2,.24,string,10,color,True)

def base(n,title,kicker,dark=True):
    sl=P.slides.add_slide(P.slide_layouts[6])
    sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG if dark else PAPER)
    rect(sl,.68,.48,.08,.19,TEAL if dark else '18867E')
    txt(sl,.91,.47,11.6,.3,kicker,11,MUTED if dark else DIM,True)
    txt(sl,.68,.97,12,1.25 if '\n' in title else .95,title,36,TEXT if dark else INK,True)
    footer(sl,n,dark)
    return sl

def footer(sl,n,dark=True):
    col=MUTED if dark else DIM
    line(sl,.7,7.02,12.65,7.02,'2A3C4D' if dark else LINE,.7)
    txt(sl,.7,7.13,6.8,.2,'ЗАГУГЛИТЬ ПРОТИВНИКА · THREAT INTELLIGENCE',9,col)
    timing=DATA['slides'][n-1]['timing']
    txt(sl,9.3,7.09,2.65,.25,timing,10,col,align=PP_ALIGN.RIGHT)
    txt(sl,12.14,7.075,.49,.29,f'{n:02}',12,col,True,align=PP_ALIGN.RIGHT)

def source(sl,n,dark=True):
    ss=DATA['slides'][n-1]['sources']
    if not ss:return
    label='  /  '.join(s['label'] for s in ss[:2])
    if len(label)>95:label='Официальная документация · полные ссылки в заметках'
    txt(sl,.73,6.68,11.9,.2,label,9,MUTED if dark else DIM,link=ss[0]['url'])

def notes(sl,n):
    item=DATA['slides'][n-1]
    note=f"СЛАЙД {n} · {item['timing']}\n\n"
    note+=item['notes']+'\n\nПЕРЕХОД: '+item['speaker_transition']
    if item['sources']:
        note+='\n\nИСТОЧНИКИ:\n'+'\n'.join(s['label']+'\n'+s['url'] for s in item['sources'])
    note+='\n\nПРОВЕРКА ДОКУМЕНТАЦИИ: 14 октября 2026 года. Текущие оценки конкретных IP/доменов не представлены.'
    if n==8:
        note+='\n\nЭто готовый запасной разбор, не скриншот интерфейса. Для live-варианта предварительно проверить публичные карточки, сохранить снимки с датой. Не открывать сам домен, не скачивать образцы, не запускать Scan/Upload/Reanalyze.'
        note+='\nhttps://www.abuseipdb.com/check/192.236.194.2\nhttps://www.virustotal.com/gui/domain/maaays.xyz'
    if n==15:
        note+='\n\nДОПОЛНИТЕЛЬНО ИЗУЧЕННЫЕ ВЫСТУПЛЕНИЯ:\n'
        note+='Aaron Shelmire, Effective Threat Intel, 2017\nhttps://speakerdeck.com/ashelmire/effective-threat-intel-sans-cti-summit-2017\n'
        note+='Alexandre Dulaunoy, The Art of Pivoting, 2025\nhttps://ail-project.org/assets/img/slides/the-art-of-pivoting.pdf\n'
        note+='Kathryn Knerler, Carson Zimmerman, Confidence in Chaos, 2022\nhttps://i.blackhat.com/EU-22/Wednesday-Briefings/EU-22-Knerler-Zimmerman-Confidence-in-Chaos.pdf\n'
        note+='Bhavna Soman, Death to the IOC, 2019\nhttps://i.blackhat.com/USA-19/Thursday/us-19-Soman-Death-To-The-IOC-Whats-Next-In-Threat-Intelligence.pdf\n'
        note+='Исходный SANS Luna Moth: прочитана карточка, запись не просмотрена, PDF не изучен.\nhttps://www.sans.org/presentations/luna-moth-a-practical-case-of-threat-intelligence-from-ioc-to-unraveling-an-attacker-infrastructure'
    if n==16:
        note+='\n\nВОПРОСЫ АУДИТОРИИ:\n1. Сколько детектов достаточно? Универсального порога нет: содержание, время, поведение и контекст.\n2. Можно закрыть при пустых сервисах? Нет: проверить разрешённость доступа и действия.\n3. Блокировать старые IP? Сначала проверить период, принадлежность и последствия.\n4. Это Luna Moth? По одному IP/инструменту определить нельзя, атрибуция требует самостоятельного обоснования.\n5. Автоматизировать? Да, с подходящими API, условиями использования и правилами передачи данных.'
    sl.notes_slide.notes_text_frame.text=note

def card(sl,x,y,w,h,head,body,dark=True,accent=TEAL,size=21):
    rect(sl,x,y,w,h,CARD if dark else WHITE,radius=True)
    rect(sl,x,y,.045,h,accent)
    txt(sl,x+.2,y+.18,w-.4,.4,head,14,accent if dark else '18867E',True)
    txt(sl,x+.2,y+.77,w-.4,h-.9,body,size,TEXT if dark else INK)

# 01 · Cover: title, one compact incident sequence.
sl=P.slides.add_slide(P.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG)
txt(sl,.72,.5,8.4,.35,'THREAT INTELLIGENCE ДЛЯ НАЧИНАЮЩИХ',12,TEAL,True)
txt(sl,.7,1.18,8.1,2.25,'Загуглить\nпротивника',62,TEXT,True)
txt(sl,.75,3.85,7.5,1.15,'Программа настоящая.\nА поддержка?',30,TEXT)
txt(sl,.75,5.42,7.9,.55,'Что именно докажет внешний поиск?',22,ORANGE)
txt(sl,.75,6.3,7.5,.35,f'{COMPANY}  /  {SPEAKER}',13,MUTED)
pill(sl,9.09,1.38,3.48,'УЧЕБНАЯ СИТУАЦИЯ',ORANGE)
for i,lab in enumerate(['Письмо о подписке','Звонок по номеру\nиз счёта','Запуск удалённой\nподдержки']):
    yy=2.12+i*1.17
    if i<2:line(sl,9.15,yy+.42,9.15,yy+1.63,'39566B',2)
    circle(sl,8.92,yy+.08,.46,TEAL)
    txt(sl,8.97,yy+.16,.35,.25,str(i+1),12,INK,True,align=PP_ALIGN.CENTER)
    txt(sl,9.63,yy,2.98,.78,lab,20,TEXT)
txt(sl,9.06,6.05,3.43,.5,'Внешний адрес — подсказка',16,MUTED)
footer(sl,1);notes(sl,1)

# 02 · Observables and contextual indicator.
sl=base(2,'Адрес — наблюдение.\nIoC — признак с основанием.','01 · ОТ СОБЫТИЯ К ОБЪЕКТУ',False)
for x,head,body in [(.7,'СЕТЬ','IP · домен · URL'),(4.82,'ФАЙЛ','Байты · имя · SHA-256'),(8.94,'КОНТЕКСТ','Письмо · телефон · сеанс')]:
    card(sl,x,2.55,3.68,1.95,head,body,False,size=22)
rect(sl,.7,5.06,11.93,.85,INK,radius=True)
txt(sl,.95,5.25,11.42,.5,'Наблюдение + роль + период + основание → IoC',25,TEXT,True,align=PP_ALIGN.CENTER)
txt(sl,.73,6.15,11.8,.43,'Совпадение — повод проверить событие.',22,DIM)
notes(sl,2)

# 03 · Five tasks, intentionally no false rigid external/internal division.
sl=base(3,'От репутации до hunting','02 · ПЯТЬ РАЗНЫХ ЗАДАЧ')
tasks=[('Репутация','Что сообщали?'),('Обогащение','Что известно?'),('Анализ TI','Что это\nобъясняет?'),('Расследование','Что произошло\nу нас?'),('Hunting','Встречалось\nли ещё?')]
for i,(a,b) in enumerate(tasks):
    xx=.72+i*2.44
    rect(sl,xx,2.65,2.16,2.5,CARD,radius=True)
    txt(sl,xx+.17,2.82,1.83,.3,f'0{i+1}',12,TEAL,True)
    txt(sl,xx+.17,3.35,1.85,.52,a,15 if i==3 else 18,TEXT,True)
    txt(sl,xx+.17,3.98,1.85,1.1,b,19,MUTED)
    if i<4:arrow(sl,xx+2.24,3.74,.14,.22)
txt(sl,.77,5.69,11.76,.75,'Поиск IP даёт подсказку.\nHunting проверяет гипотезу в событиях.',24,LIME)
notes(sl,3)

# 04 · Four object -> question -> source rows.
sl=base(4,'Сначала вопрос. Потом источник.','03 · ВЫБИРАЕМ, ЧТО ИСКАТЬ',False)
rows=DATA['slides'][3]['screen']['rows']
for i,r in enumerate(rows):
    yy=2.12+i*.92
    rect(sl,.7,yy,11.95,.78,WHITE,radius=True)
    txt(sl,.95,yy+.18,2.22,.5,r['object'].replace('Сеанс / процесс','Сеанс'),20,INK,True)
    txt(sl,3.27,yy+.18,5.32,.44,r['question'],20,INK)
    txt(sl,9.05,yy+.18,3.28,.45,r['source'].replace('VirusTotal · нужный объект','VirusTotal'),18,'18867E',True)
txt(sl,.78,6.07,11.79,.43,'Домен не заменяет полный URL. Имя файла не заменяет хеш.',21,DIM)
source(sl,4,False);notes(sl,4)

# 05 · No scores: reading questions and interpretation.
sl=base(5,'AbuseIPDB: читаем историю','04 · СООБЩЕНИЯ ОБ IP')
for x,y,a,b in [(.72,2.22,'ПЕРИОД','Какое окно\nпроверяем?'),(4.68,2.22,'ДАТЫ','Когда\nсообщали?'),(.72,4.1,'КАТЕГОРИИ','Что наблюдали?'),(4.68,4.1,'РЕПОРТЁРЫ','Кто и сколько раз?')]:
    card(sl,x,y,3.68,1.64,a,b,size=20)
txt(sl,8.82,2.5,3.7,1.12,'≠',78,LIME,True)
txt(sl,8.89,3.76,3.58,.95,'Оценка сервиса\nне равна вероятности\nнашего инцидента',23,TEXT,True)
txt(sl,8.91,5.12,3.6,.54,'Комментарии помогают\nпонять содержание жалобы.',16,MUTED)
txt(sl,.78,6.05,11.8,.5,'Количество сообщений ≠ количество атак. Нет жалоб ≠ безопасно.',20,ORANGE)
source(sl,5);notes(sl,5)

# 06 · VT object types and normalized verdicts.
sl=base(6,'VirusTotal: читаем весь отчёт','05 · ТИП ОБЪЕКТА И ВРЕМЯ',False)
for i,l in enumerate(['IP','ДОМЕН','URL','ФАЙЛ']):
    pill(sl,.75+i*1.79,2.05,1.5,l,'18867E',False)
txt(sl,.76,2.91,6.25,1.37,'Доля детектов\n≠ вероятность атаки',32,INK,True)
txt(sl,.78,4.67,5.9,1.0,'Читать дальше:\nдату · тип · поведение · связи',22,DIM)
verdicts=[('malicious','Считают вредоносным','C64646'),('undetected','Нет детекта / мнения',DIM),('harmless','Считают невредоносным','18867E'),('timeout / failure','Результат не получен','A86D2D')]
for i,(head,body,col) in enumerate(verdicts):
    yy=2.7+i*.73
    line(sl,7.7,yy+.11,7.7,yy+.57,col,3)
    txt(sl,7.89,yy,4.61,.31,head,18,col,True,font=MONO)
    txt(sl,7.89,yy+.35,4.62,.31,body,17,INK)
rect(sl,.7,6.03,11.94,.54,INK,radius=True)
txt(sl,.95,6.15,11.39,.3,'Легитимная программа ≠ разрешённая сессия.',21,TEXT,True)
source(sl,6,False);notes(sl,6)

# 07 · Disclosure and three distinct operations.
sl=base(7,'Что раскрывает внешний запрос?','06 · ПЕРЕД ПОИСКОМ')
for x,a,b in [(.72,'ОТЧЁТ ПО ХЕШУ','Запрос известного\nотпечатка'),(4.83,'UPLOAD','Содержимое\nфайла'),(8.94,'REANALYZE','Новый анализ\nобъекта')]:
    card(sl,x,2.25,3.69,1.9,a,b,size=22)
rect(sl,.72,4.64,11.92,1.2,'3E302A',radius=True)
txt(sl,.99,4.84,11.35,.79,'VT предупреждает об анализе и включении запрошенных\nи переданных IoC в набор данных сообщества.',23,ORANGE,True)
txt(sl,.79,6.16,11.69,.4,'Приватные ссылки, токены и документы — в разрешённом контуре.',20,MUTED)
source(sl,7);notes(sl,7)

# 08 · Ready static fallback: actual historical published pair, no fake report UI.
sl=base(8,'Одна пара — два вопроса','07 · РАЗБОР ПУБЛИКАЦИИ')
pill(sl,.75,2.13,4.02,'SYGNIA · ИСТОРИЧЕСКИЙ ОБЪЕКТ · 2022',LIME)
rect(sl,.72,2.78,5.2,.76,CARD,radius=True)
txt(sl,.98,2.95,4.7,.5,'maaays[.]xyz',29,TEXT,True,font=MONO)
txt(sl,6.08,2.88,.75,.57,'↔',35,TEAL,True,align=PP_ALIGN.CENTER)
rect(sl,7.02,2.78,5.61,.76,CARD,radius=True)
txt(sl,7.24,2.95,5.16,.5,'192[.]236[.]194[.]2',27,TEXT,True,font=MONO)
txt(sl,.77,3.78,11.75,.37,'Роль в таблице: Exfiltration Server — сервер для выноса данных.',20,MUTED)
card(sl,.73,4.39,5.72,1.61,'ABUSEIPDB','Что сообщали об IP — и когда?',size=21)
card(sl,6.87,4.39,5.75,1.61,'VIRUSTOTAL','Что известно о домене\nи датах связей?',size=21)
txt(sl,.8,6.22,11.69,.38,'Текущих оценок нет. К учебному компьютеру эта пара не привязана.',20,ORANGE)
source(sl,8);notes(sl,8)

# 09 · Published attack chain, explicit campaign scope.
sl=base(9,'Luna Moth: цепочка обмана','08 · РЕАЛЬНЫЙ КЕЙС')
pill(sl,.77,2.05,3.54,'ЦЕПОЧКА ПО SYGNIA · 2022',LIME)
chain=[('Счёт','PDF о подписке'),('Звонок','Жертва звонит сама'),('Поддержка','Zoho Assist'),('Управление','Atera'),('Вынос данных','Rclone')]
for i,(a,b) in enumerate(chain):
    xx=.73+i*2.43
    circle(sl,xx+.63,2.53,.74,TEAL if i<4 else ORANGE)
    txt(sl,xx+.73,2.72,.54,.32,str(i+1),17,INK,True,align=PP_ALIGN.CENTER)
    if i<4:arrow(sl,xx+1.86,2.8,.3,.27)
    txt(sl,xx,3.56,2.15,.65,a,18,TEXT,True,align=PP_ALIGN.CENTER)
    txt(sl,xx,4.36,2.15,.75,b,17,MUTED,align=PP_ALIGN.CENTER)
txt(sl,.8,5.55,11.69,.47,'Знакомый инструмент. Посторонний оператор.',27,LIME,True)
txt(sl,.8,6.13,11.75,.4,'Кампания Unit 42: хищение и вымогательство без шифрования.',19,MUTED)
source(sl,9);notes(sl,9)

# 10 · Common hosting example, schematic and explicitly not discovered assets.
sl=base(10,'У связи есть время и основание','09 · ПЕРЕХОД К КОНТЕКСТУ',False)
for i,(a,b) in enumerate([('Источник','Откуда связь?'),('Время','Когда наблюдали?'),('Характер','Что связывает объекты?')]):
    x=.76+i*4.06
    txt(sl,x,2.13,3.62,.42,a,20,'18867E',True)
    txt(sl,x,2.68,3.61,.45,b,20,INK)
line(sl,3.15,4.12,6.21,4.12,'7795A9',2)
for yy in [3.55,4.18,4.81]:line(sl,6.42,4.12,9.52,yy,'7795A9',1.5,dash=True)
rect(sl,.91,3.69,2.68,.86,WHITE,LINE,True);txt(sl,1.17,3.95,2.16,.41,'Домен A',22,INK,True)
rect(sl,4.79,3.63,2.99,1.02,INK,None,True);txt(sl,5.08,3.82,2.42,.66,'Общий IP\nразные клиенты',19,TEXT,True,align=PP_ALIGN.CENTER)
for yy,l in [(3.29,'Сайт B'),(3.92,'Сайт C'),(4.55,'Сайт D')]:
    rect(sl,9.48,yy,2.97,.5,WHITE,LINE,True);txt(sl,9.75,yy+.1,2.4,.31,l,17,INK)
txt(sl,.98,4.95,10.98,.29,'Учебная схема общего хостинга — не найденная инфраструктура.',14,DIM)
txt(sl,.81,5.68,11.67,.63,'Кто разрешил доступ — и что сделал оператор?',27,INK,True)
source(sl,10,False);notes(sl,10)

# 11 · Internal evidence collection plan, not a fabricated timeline.
sl=base(11,'Что произошло у нас?','10 · ВНУТРЕННИЕ ИСТОЧНИКИ')
rows=[('Почта и сотрудник','Когда и откуда начался контакт?'),('Устройство и процесс','Что запустили?'),('Сессия поддержки','Кто подключался? Было ли разрешение?'),('Сеть и файлы','Какие обращения и передача видны?')]
for i,(a,b) in enumerate(rows):
    yy=2.17+i*.88
    txt(sl,.82,yy,4.45,.41,a,20,TEAL,True)
    txt(sl,5.49,yy,6.9,.5,b,22,TEXT)
    line(sl,.82,yy+.63,12.52,yy+.63,'2A3C4D',.8)
txt(sl,.81,5.97,11.71,.39,'Сопоставить: устройство · время · пользователь · процесс / сеанс',19,MUTED)
txt(sl,.81,6.5,11.71,.36,'DNS-запрос ≠ соединение. Соединение ≠ кража.',21,ORANGE,True)
notes(sl,11)

# 12 · A complete training record, explicit unperformed verification.
sl=base(12,'Вывод, который можно проверить','11 · УЧЕБНАЯ ЗАПИСЬ',False)
pill(sl,.77,2.04,4.62,'УЧЕБНЫЙ ПРИМЕР · ПРОВЕРКА НЕ ВЫПОЛНЕНА','B17438',False)
for i,r in enumerate(DATA['slides'][11]['screen']['note_rows']):
    yy=2.71+i*.63
    txt(sl,.9,yy,2.01,.35,r['label'],19,'18867E',True)
    txt(sl,3.05,yy,9.39,.41,r['text'],21,INK)
    line(sl,.9,yy+.48,12.48,yy+.48,LINE,.7)
txt(sl,.87,6.17,11.66,.44,'Сохраняем источник, дату и пределы вывода.',22,DIM)
notes(sl,12)

# 13 · Hypothesis and feasible first action, no fake success.
sl=base(13,'Hunting начинается с гипотезы','12 · ПОИСК В ИСТОРИИ')
pill(sl,.77,2.06,3.78,'14 ДНЕЙ — УСЛОВИЕ УПРАЖНЕНИЯ',ORANGE)
txt(sl,.8,2.78,11.7,.46,'Возможное неразрешённое использование:',22,MUTED)
txt(sl,.8,3.38,11.7,.65,'Удалённая поддержка → передача файлов',31,TEXT,True)
for i,(a,b) in enumerate([('НАЙТИ','Запуски и сеансы'),('ПРОВЕРИТЬ','Последующие действия'),('СВЕРИТЬ','Заявки и операторов')]):
    xx=.77+i*4.09
    card(sl,xx,4.54,3.72,1.37,a,b,size=19)
txt(sl,.82,6.16,11.66,.42,'До поиска: охват устройств, срок хранения и время событий.',20,MUTED)
notes(sl,13)

# 14 · Close on actionable method.
sl=P.slides.add_slide(P.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG)
txt(sl,.76,.59,10.3,.29,'THREAT INTELLIGENCE В РАБОТЕ',12,TEAL,True)
txt(sl,.73,1.43,8.45,2.76,'Один объект.\nОдин вопрос.\nСледующий шаг.',48,TEXT,True)
txt(sl,.81,4.87,7.49,1.06,'Проверить разрешённость сеанса\nи действия оператора.',25,LIME)
txt(sl,.81,6.36,7.47,.35,'Спасибо! Вопросы?',18,MUTED)
for i,l in enumerate(['Объект','Вопрос','Источник','Проверка','Действие']):
    yy=1.66+i*.89
    circle(sl,9.37,yy+.04,.16,TEAL)
    if i<4:line(sl,9.45,yy+.21,9.45,yy+.93,'39566B',1.5)
    txt(sl,9.9,yy,2.8,.51,l,25,TEXT if i<4 else LIME,True)
footer(sl,14);notes(sl,14)

# 15 · Clickable primary sources; no raw URL noise on screen.
sl=base(15,'Источники, на которые можно опереться','ПРИЛОЖЕНИЕ · ОФИЦИАЛЬНЫЕ МАТЕРИАЛЫ',False)
groups=[('КЕЙС',[("Sygnia · Luna Moth, 2022",'https://www.sygnia.co/blog/luna-moth-false-subscription-scams/'),('Unit 42 · Callback Phishing, 2022','https://unit42.paloaltonetworks.com/luna-moth-callback-phishing/')]),('РЕПУТАЦИЯ IP',[('AbuseIPDB · CHECK','https://docs.abuseipdb.com/#check-endpoint'),('AbuseIPDB · FAQ','https://www.abuseipdb.com/faq.html')]),('ОТЧЁТЫ И ДАННЫЕ',[('VT · Analysis object','https://docs.virustotal.com/reference/analyses-object'),('VT · Searching','https://docs.virustotal.com/docs/searching'),('VT · Upload / disclosure','https://docs.virustotal.com/reference/files-scan')])]
for i,(head,links) in enumerate(groups):
    xx=.77+i*4.08
    card(sl,xx,2.37,3.73,3.2,head,'',False)
    for j,(lab,url) in enumerate(links):
        txt(sl,xx+.23,3.31+j*.71,3.27,.61,lab,19,INK,link=url)
txt(sl,.87,6.03,11.55,.72,'Проверено 14 октября 2026 года.\nПолные ссылки и основания — в заметках слайдов.',21,DIM)
notes(sl,15)

# 16 · Take-away card.
sl=base(16,'Сохранить себе: проверка объекта','ПРИЛОЖЕНИЕ · ПАМЯТКА АНАЛИТИКУ')
check=['Зафиксировать событие\nи время','Выбрать объект\nи вопрос','Проверить допустимость\nвнешнего запроса','Прочитать содержание,\nисточник и даты','Сопоставить\nс внутренними данными','Записать вывод,\nпробелы и действие']
for i,l in enumerate(check):
    col=0 if i<3 else 1; row=i%3
    xx=.8+col*6.08; yy=2.34+row*1.04
    circle(sl,xx,yy+.05,.46,TEAL)
    txt(sl,xx+.065,yy+.15,.32,.24,str(i+1),12,INK,True,align=PP_ALIGN.CENTER)
    txt(sl,xx+.75,yy,4.86,.89,l,22,TEXT)
rect(sl,.79,5.94,11.78,.59,CARD,radius=True)
txt(sl,1.02,6.08,11.31,.33,'Событие → свидетельство → проверка → вывод → действие',21,LIME,True)
notes(sl,16)

assert len(P.slides)==16
for sn,st,x,y,w,h,size in TRACK:
    assert x>=0 and y>=0 and x+w<=13.334 and y+h<=7.5,(sn,st,'out of bounds')
dest=ROOT/'ti_presentation.pptx';P.save(dest)
out=[]
for i,s in enumerate(P.slides,1):
    out.append(f'# Слайд {i}\n\n'+s.notes_slide.notes_text_frame.text)
(ROOT/'ti_presentation_speaker_notes.md').write_text('\n\n'.join(out))
print(json.dumps({'pptx':str(dest),'bytes':dest.stat().st_size,'slides':len(P.slides),'text_boxes':len(TRACK),'notes_slides':sum(bool(s.notes_slide.notes_text_frame.text) for s in P.slides)},ensure_ascii=False))
